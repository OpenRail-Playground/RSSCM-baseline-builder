"""Ladeliste parser — extract metadata from vendor Excel/CSV delivery."""
import io
import zipfile
from pathlib import Path

import openpyxl


def parse_ladeliste(path: str) -> list[dict]:
    """Parse a vendor Ladeliste (xlsx, csv, or zip of xlsx) into normalized software items."""
    p = Path(path)
    if p.suffix.lower() == ".zip":
        return _parse_zip(p)
    elif p.suffix.lower() in (".xlsx", ".xls"):
        return _parse_excel(p)
    elif p.suffix.lower() == ".csv":
        return _parse_csv(p)
    raise ValueError(f"Unsupported Ladeliste format: {p.suffix}")


def _parse_zip(zip_path: Path) -> list[dict]:
    items = []
    with zipfile.ZipFile(zip_path, "r") as zf:
        for name in zf.namelist():
            if name.endswith(".xlsx"):
                with zf.open(name) as f:
                    items.extend(_parse_excel_bytes(f.read(), source=name))
    return items


def _parse_excel(path: Path) -> list[dict]:
    return _parse_excel_bytes(path.read_bytes(), source=str(path))


def _parse_excel_bytes(data: bytes, source: str = "") -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    # Prefer "Flat Data" sheet if present
    sheet_name = "Flat Data" if "Flat Data" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header = [str(h).strip().lower() if h else "" for h in rows[0]]
    items = []

    for row in rows[1:]:
        if not row or not any(row):
            continue
        record = {header[i]: (str(row[i]).strip() if row[i] else "") for i in range(min(len(header), len(row)))}
        item = _map_to_model(record, source)
        if item:
            items.append(item)

    wb.close()
    return items


def _map_to_model(record: dict, source: str) -> dict | None:
    """Map a flat-data row to the RSSCM data model."""
    sw_type = record.get("software type", "")
    version = record.get("sw version", "")
    if not sw_type or not version:
        return None

    return {
        "subsystem": record.get("subsystem", ""),
        "system_component": record.get("system component", ""),
        "hardware": record.get("hardware (board)", ""),
        "part_number": record.get("part number", ""),
        "software_type": _normalize_software_type(sw_type),
        "name": sw_type,
        "version": version,
        "release": record.get("release", ""),
        "file_artifact": record.get("file / artifact", ""),
        "tool": record.get("tool", ""),
        "source": source,
    }


def _normalize_software_type(raw: str) -> str:
    """Map vendor software type strings to RSSCM SoftwareType choices."""
    lower = raw.lower()
    if "operating system" in lower or "os" == lower:
        return "Operating System"
    if "application" in lower or "appl" in lower:
        return "Application"
    if "bootloader" in lower or "boot" in lower:
        return "Bootloader"
    if "config" in lower or "parameter" in lower:
        return "Configuration"
    if "firmware" in lower or "fw" in lower:
        return "Firmware"
    if "driver" in lower:
        return "Drivers"
    return "Other"


def _parse_csv(path: Path) -> list[dict]:
    import csv

    items = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            record = {k.strip().lower(): v.strip() for k, v in row.items() if k}
            item = _map_to_model(record, source=str(path))
            if item:
                items.append(item)
    return items
